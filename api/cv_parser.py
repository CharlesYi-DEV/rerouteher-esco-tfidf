from __future__ import annotations

import io
import re
import zipfile
from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from pypdf import PdfReader


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}
MONTH_PATTERN = "|".join(sorted(MONTHS, key=len, reverse=True))
DATE_TOKEN = rf"(?:(?:{MONTH_PATTERN})[\s./-]+(?:19|20)\d{{2}}|(?:0?[1-9]|1[0-2])[/.-](?:19|20)\d{{2}}|(?:19|20)\d{{2}})"
DATE_RANGE_RE = re.compile(
    rf"(?P<start>{DATE_TOKEN})\s*(?:-|–|—|to|until)\s*(?P<end>{DATE_TOKEN}|present|current|now)",
    re.IGNORECASE,
)

EXPERIENCE_HEADINGS = {
    "experience",
    "work experience",
    "professional experience",
    "employment history",
    "work history",
    "career history",
}
SKILLS_HEADINGS = {
    "skills",
    "technical skills",
    "core skills",
    "core competencies",
    "core competences",
    "professional skills",
    "tools and technologies",
    "technologies",
}
OTHER_HEADINGS = {
    "profile",
    "summary",
    "professional summary",
    "objective",
    "education",
    "qualifications",
    "certifications",
    "projects",
    "languages",
    "awards",
    "references",
    "interests",
    "volunteering",
}
ALL_HEADINGS = EXPERIENCE_HEADINGS | SKILLS_HEADINGS | OTHER_HEADINGS

ROLE_WORDS = {
    "accountant",
    "advisor",
    "associate",
    "administrator",
    "analyst",
    "architect",
    "assistant",
    "cashier",
    "chef",
    "clerk",
    "consultant",
    "coordinator",
    "developer",
    "designer",
    "director",
    "editor",
    "engineer",
    "executive",
    "founder",
    "head",
    "intern",
    "manager",
    "nurse",
    "officer",
    "operator",
    "programmer",
    "recruiter",
    "representative",
    "salesperson",
    "scientist",
    "specialist",
    "supervisor",
    "teacher",
    "technician",
    "writer",
    "worker",
    "waiter",
    "lead",
}
COMPANY_WORDS = {"berhad", "bhd", "company", "corp", "corporation", "inc", "limited", "llc", "ltd", "plc", "sdn"}


class CvParseError(ValueError):
    pass


@dataclass
class ParsedCv:
    job_title: str
    skills: list[str]
    work_length_years: float | None
    total_experience_years: float | None
    job_title_source: str
    skills_source: str
    work_length_method: str
    employment_date_ranges_found: int
    text_characters: int
    text_preview: str
    warnings: list[str]

    def to_dict(self) -> dict:
        return asdict(self)


def _clean_lines(text: str) -> list[str]:
    lines = []
    for raw in text.replace("\x00", "").replace("\r", "\n").split("\n"):
        line = re.sub(r"[\t\u00a0 ]+", " ", raw).strip()
        if line:
            lines.append(line)
    return lines


def _heading(line: str) -> str | None:
    candidate = re.sub(r"[:\-–—]+$", "", line).strip().lower()
    candidate = re.sub(r"\s+", " ", candidate)
    return candidate if candidate in ALL_HEADINGS else None


def _extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise CvParseError("Supported CV formats are PDF, DOCX, and TXT.")
    if not content:
        raise CvParseError("The uploaded CV is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise CvParseError("The uploaded CV exceeds the 10 MB test limit.")

    try:
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                try:
                    unlocked = reader.decrypt("")
                except Exception as exc:  # pragma: no cover - library-specific
                    raise CvParseError("Encrypted PDFs are not supported.") from exc
                if not unlocked:
                    raise CvParseError("Encrypted PDFs are not supported.")
            if len(reader.pages) > 50:
                raise CvParseError("The PDF exceeds the 50-page test limit.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 5_000 or sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise CvParseError("The DOCX expands beyond the safe internal-test limit.")
            document = Document(io.BytesIO(content))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    blocks.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(blocks)
        else:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("latin-1")
    except CvParseError:
        raise
    except Exception as exc:
        raise CvParseError(f"The {extension[1:].upper()} file could not be read.") from exc

    text = text.strip()
    if len(text) < 20:
        raise CvParseError("No usable CV text was extracted. Scanned PDFs require OCR before this test.")
    return text


def _section(lines: list[str], headings: set[str]) -> tuple[list[str], str | None]:
    for index, line in enumerate(lines):
        heading = _heading(line)
        if heading in headings:
            collected = []
            for following in lines[index + 1 :]:
                if _heading(following):
                    break
                collected.append(following)
            return collected, heading
    return [], None


def _role_score(value: str) -> int:
    normalized = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
    words = normalized.split()
    if not words or len(words) > 12 or len(value) > 120:
        return -20
    if "@" in value or "http" in normalized or re.search(r"\+?\d[\d ()-]{7,}", value):
        return -20
    score = 0
    score += 4 * len(set(words) & ROLE_WORDS)
    score -= 3 * len(set(words) & COMPANY_WORDS)
    if 2 <= len(words) <= 6:
        score += 2
    if re.search(r"\b(responsible|duties|achieved|managed|developed|worked)\b", normalized):
        score -= 5
    return score


def _best_title_fragment(value: str) -> tuple[str, int] | None:
    cleaned = DATE_RANGE_RE.sub("", value)
    cleaned = re.sub(r"^[|,;:()\-–—\s]+|[|,;:()\-–—\s]+$", "", cleaned)
    if not cleaned or _heading(cleaned):
        return None
    pieces = re.split(r"\s+(?:at|@)\s+|\s*[|•·,]\s*", cleaned, flags=re.IGNORECASE)
    scored = [(piece.strip(), _role_score(piece.strip())) for piece in pieces if piece.strip()]
    if not scored:
        return None
    candidate, score = max(scored, key=lambda item: item[1])
    return (candidate, score) if score >= 4 else None


def _parse_date_token(token: str, *, is_end: bool) -> int:
    value = token.lower().strip().replace(".", " ")
    today = date.today()
    if value in {"present", "current", "now"}:
        return today.year * 12 + today.month - 1
    numeric = re.search(r"\b(0?[1-9]|1[0-2])[/\-]((?:19|20)\d{2})\b", value)
    if numeric:
        month, year = int(numeric.group(1)), int(numeric.group(2))
        return year * 12 + month - 1
    year_match = re.search(r"\b((?:19|20)\d{2})\b", value)
    if not year_match:
        raise CvParseError("A detected employment date could not be interpreted.")
    year = int(year_match.group(1))
    month = next((number for name, number in MONTHS.items() if re.search(rf"\b{name}\b", value)), None)
    month = month or (12 if is_end else 1)
    return year * 12 + month - 1


def _date_ranges(text: str) -> list[tuple[int, int]]:
    ranges = []
    for match in DATE_RANGE_RE.finditer(text):
        start = _parse_date_token(match.group("start"), is_end=False)
        end = _parse_date_token(match.group("end"), is_end=True)
        if end >= start and end - start <= 80 * 12:
            ranges.append((start, end))
    return ranges


def _merged_years(ranges: list[tuple[int, int]]) -> float | None:
    if not ranges:
        return None
    intervals = sorted((start, end + 1) for start, end in ranges)
    merged: list[list[int]] = []
    for start, end in intervals:
        if not merged or start > merged[-1][1]:
            merged.append([start, end])
        else:
            merged[-1][1] = max(merged[-1][1], end)
    months = sum(end - start for start, end in merged)
    return round(min(months / 12, 80), 2)


def _extract_job_title(
    lines: list[str], experience_lines: list[str]
) -> tuple[str, str, tuple[int, int] | None]:
    scope = experience_lines or lines
    candidates: list[tuple[int, int, int, str, str]] = []
    for index, line in enumerate(scope):
        matches = list(DATE_RANGE_RE.finditer(line))
        if not matches:
            continue
        latest_match = max(matches, key=lambda match: _parse_date_token(match.group("end"), is_end=True))
        start = _parse_date_token(latest_match.group("start"), is_end=False)
        end = _parse_date_token(latest_match.group("end"), is_end=True)
        nearby = [(index, line, "same line")]
        for offset, label in [(-1, "line before date"), (-2, "two lines before date"), (1, "line after date")]:
            position = index + offset
            if 0 <= position < len(scope):
                nearby.append((position, scope[position], label))
        for _, value, label in nearby:
            fragment = _best_title_fragment(value)
            if fragment:
                title, score = fragment
                candidates.append((end, score, start, title, label))

    if candidates:
        candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
        end, _, start, title, source = candidates[0]
        return title, f"latest employment date range; {source}", (start, end)

    fallback = []
    for index, line in enumerate(scope[:20]):
        fragment = _best_title_fragment(line)
        if fragment:
            title, score = fragment
            fallback.append((score, -index, title))
    if fallback:
        _, _, title = max(fallback)
        return title, "role-like line near experience section", None
    raise CvParseError("A job title could not be extracted from the CV experience section.")


def _split_skills(values: Iterable[str]) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for value in values:
        for candidate in re.split(r"[,;|•·\n]+", value):
            candidate = re.sub(r"^[\-–—*✓►▪●\s]+", "", candidate).strip(" .:")
            if ":" in candidate:
                prefix, suffix = candidate.split(":", 1)
                if 1 <= len(prefix.split()) <= 4:
                    candidate = suffix.strip() or prefix.strip()
            normalized = re.sub(r"[^a-z0-9+#.]+", " ", candidate.lower()).strip()
            if not normalized or normalized in seen:
                continue
            if len(candidate) > 80 or len(candidate.split()) > 10 or DATE_RANGE_RE.search(candidate):
                continue
            seen.add(normalized)
            output.append(candidate)
            if len(output) >= 40:
                return output
    return output


def _extract_skills(lines: list[str], text: str, official_skill_groups: Iterable[str]) -> tuple[list[str], str]:
    section, heading = _section(lines, SKILLS_HEADINGS)
    skills = _split_skills(section[:20]) if section else []
    sources = [f"{heading} section"] if skills and heading else []

    if not skills:
        labelled = []
        for line in lines:
            match = re.match(r"^(?:tools|technologies|skills|proficient in)\s*[:\-]\s*(.+)$", line, re.IGNORECASE)
            if match:
                labelled.append(match.group(1))
        skills = _split_skills(labelled)
        if skills:
            sources.append("labelled skill lines")

    normalized_text = re.sub(r"[^a-z0-9]+", " ", text.lower())
    official_matches = []
    for skill in sorted(set(official_skill_groups), key=len, reverse=True):
        normalized = re.sub(r"[^a-z0-9]+", " ", skill.lower()).strip()
        if len(normalized.split()) >= 2 and normalized in normalized_text:
            official_matches.append(skill)
        if len(official_matches) >= 20:
            break
    if official_matches:
        skills = _split_skills([*skills, *official_matches])
        sources.append("official ESCO skill-group phrase matches")

    return skills, " + ".join(sources) if sources else "no explicit skills extracted"


def parse_cv(filename: str, content: bytes, official_skill_groups: Iterable[str] = ()) -> ParsedCv:
    text = _extract_text(filename, content)
    warnings: list[str] = []
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        warnings.append("Extracted text was truncated to 100,000 characters.")
    lines = _clean_lines(text)
    experience_lines, experience_heading = _section(lines, EXPERIENCE_HEADINGS)
    if not experience_heading:
        warnings.append("No standard experience heading was found; title extraction used the full CV.")

    job_title, job_title_source, title_range = _extract_job_title(lines, experience_lines)
    skills, skills_source = _extract_skills(lines, text, official_skill_groups)
    if not skills:
        warnings.append("No explicit skills were extracted; both models will rely mainly on the job title.")

    duration_text = "\n".join(experience_lines) if experience_lines else text
    ranges = _date_ranges(duration_text)
    if not ranges and experience_lines:
        ranges = _date_ranges(text)
        if ranges:
            warnings.append("No experience dates were found in the section; duration used all CV date ranges.")
    total_experience = _merged_years(ranges)
    if title_range:
        work_length = round(min((title_range[1] - title_range[0] + 1) / 12, 80), 2)
        work_length_method = "latest role employment date range"
    elif ranges:
        latest_range = max(ranges, key=lambda item: item[1])
        work_length = round(min((latest_range[1] - latest_range[0] + 1) / 12, 80), 2)
        work_length_method = "latest detected employment date range"
        warnings.append("The latest date range could not be tied directly to the extracted title.")
    else:
        work_length = None
        work_length_method = "not available"
        warnings.append("No employment date range was found; work length was omitted from model input.")

    preview = "\n".join(lines)[:1500]
    return ParsedCv(
        job_title=job_title,
        skills=skills,
        work_length_years=work_length,
        total_experience_years=total_experience,
        job_title_source=job_title_source,
        skills_source=skills_source,
        work_length_method=work_length_method,
        employment_date_ranges_found=len(ranges),
        text_characters=len(text),
        text_preview=preview,
        warnings=warnings,
    )
